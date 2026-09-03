"""
Regenera data.js (el buscador) a partir de actas_index.xlsx.

PARA QUÉ SIRVE
--------------
Cuando corriges algo a mano en el Excel — arreglas un nombre, una
cédula, o borras el "SÍ" de la columna Revisar porque ya lo
verificaste — esa corrección NO se refleja sola en el buscador. El
Excel y data.js quedan como dos copias independientes desde el
momento en que procesar_actas.py los generó la primera vez.

Este script hace que el Excel sea la única fuente de verdad después
de que lo revisas: lee actas_index.xlsx tal como quedó (con tus
correcciones) y regenera data.js desde ahí — así el buscador queda
sincronizado con un solo comando, sin tocar el HTML ni editar JSON
a mano.

FLUJO RECOMENDADO
------------------
1. Corriges lo que haga falta directo en actas_index.xlsx (nombres,
   cédulas, sedes), y borras el "SÍ" de la columna Revisar en las
   filas que ya verificaste a mano.
2. Guardas el Excel.
3. Corres: python actualizar_buscador_desde_excel.py
4. buscador_actas.html ya queda al día — no hay que tocarlo.

USO
---
python actualizar_buscador_desde_excel.py
"""

import json
import re
from pathlib import Path
from openpyxl import load_workbook

SALIDA_EXCEL = Path("./actas_index.xlsx")
SALIDA_DATAJS = Path("./data.js")


def limpiar_nombre(texto: str) -> str:
    """Deja el nombre solo con letras y espacios simples, con
    Mayúscula Inicial en cada palabra — igual que hace
    procesar_actas.py, para que una corrección a mano en el Excel
    quede formateada igual sin que tengas que escribirla perfecta."""
    solo_letras = re.sub(r"[^A-Za-zÁÉÍÓÚÑÜáéíóúñü\s]", "", texto or "")
    un_espacio = re.sub(r"\s+", " ", solo_letras).strip()
    return un_espacio.title()


def celda_a_texto(valor) -> str:
    """Convierte el valor de una celda a texto plano, cuidando que
    Excel no haya convertido una cédula larga en número (lo que le
    metería un '.0' al final)."""
    if valor is None:
        return ""
    if isinstance(valor, float):
        return str(int(valor))
    return str(valor).strip()


def main():
    if not SALIDA_EXCEL.exists():
        print(f"No encontré {SALIDA_EXCEL.resolve()} — corre procesar_actas.py primero.")
        return

    wb = load_workbook(SALIDA_EXCEL)
    ws = wb.active

    registros = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        cedula, nombre, sede, archivo, revisar_txt = fila
        cedula = re.sub(r"[^\d]", "", celda_a_texto(cedula))
        nombre = limpiar_nombre(celda_a_texto(nombre))
        sede = celda_a_texto(sede)
        archivo = celda_a_texto(archivo)

        if not any([cedula, nombre, sede, archivo]):
            continue  # fila vacía

        revisar = celda_a_texto(revisar_txt).upper() == "SÍ"

        registros.append({
            "cedula": cedula,
            "nombre": nombre,
            "sede": sede,
            "archivo": archivo,
            "revisar": revisar,
        })

    contenido = "const ACTAS = " + json.dumps(registros, ensure_ascii=False, indent=2) + ";\n"
    SALIDA_DATAJS.write_text(contenido, encoding="utf-8")

    n_revisar = sum(1 for r in registros if r["revisar"])
    print(f"data.js actualizado: {len(registros)} personas, {n_revisar} todavía marcadas 'revisar'.")


if __name__ == "__main__":
    main()
